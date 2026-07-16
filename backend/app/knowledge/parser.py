"""知识库文档解析 + 切块(V2 §3)。

按文件扩展名分派解析器:txt/md 直读;PDF 用 PyMuPDF 提取文字;docx 用 python-docx;
xlsx 用 openpyxl 按行。解析库装不上/格式不支持时抛明确错误(上层捕获标 failed)。
切块:中文友好(按段落 + 长度滑窗),自写 ~30 行,不引 LlamaIndex。
"""
from __future__ import annotations

import logging

logger = logging.getLogger("knowledge.parser")

# 支持的扩展名(小写,含点)
SUPPORTED_EXTS = {".txt", ".md", ".markdown", ".pdf", ".docx", ".xlsx"}


def extract_text(filename: str, raw: bytes) -> str:
    """从原始字节提取纯文本。按扩展名分派。"""
    name = filename.lower()
    if name.endswith((".txt", ".md", ".markdown")):
        # 文本类:尝试 utf-8,失败回落 gbk(中文常见)
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("gbk", errors="ignore")
    if name.endswith(".pdf"):
        return _extract_pdf(raw)
    if name.endswith(".docx"):
        return _extract_docx(raw)
    if name.endswith(".xlsx"):
        return _extract_xlsx(raw)
    raise ValueError(f"不支持的文件类型: {filename}(支持 {sorted(SUPPORTED_EXTS)})")


def _extract_pdf(raw: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ValueError("PDF 解析未启用(PyMuPDF 未安装)") from e
    text_parts: list[str] = []
    with fitz.open(stream=raw, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n\n".join(text_parts).strip()


def _extract_docx(raw: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise ValueError("Word 解析未启用(python-docx 未安装)") from e
    import io
    doc = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(raw: bytes) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise ValueError("Excel 解析未启用(openpyxl 未安装)") from e
    import io
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)


def chunk_text(text: str, size: int = 500, overlap: int = 50) -> list[str]:
    """中文友好的切块:先按段落/换行切,再按 size 滑窗合并,带 overlap。

    size/overlap 按字符计(中文场景 token≈字符,粗略够用)。
    """
    text = text.strip()
    if not text:
        return []
    # 按空行/换行拆成段落
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        # 段落本身超长 → 按 size 硬切
        if len(para) > size:
            if buf:
                chunks.append(buf)
                buf = ""
            for i in range(0, len(para), size - overlap):
                chunks.append(para[i:i + size])
                if i + size >= len(para):
                    break
            continue
        # 累积到 buf,超 size 则落盘
        if len(buf) + len(para) + 1 > size and buf:
            chunks.append(buf)
            # overlap:保留 buf 末尾 overlap 字符作下一块开头
            buf = buf[-overlap:] + "\n" + para if overlap else para
        else:
            buf = (buf + "\n" + para) if buf else para
    if buf:
        chunks.append(buf)
    return [c for c in chunks if c.strip()]
